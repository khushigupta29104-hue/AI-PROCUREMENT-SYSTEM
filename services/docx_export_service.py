"""
docx_export_service.py
========================
Fills a proposal template with resolved field values, replacing every
{{placeholder_token}} in paragraphs AND table cells with the final
value (which may have been edited by the user in the Proposal Preview
step before this export runs). Saves the result to
generated_proposals/docx/.
"""

import os
from datetime import datetime
from typing import Dict

from docx import Document

from utils.config_loader import get_field_mapping, get_proposal_config, get_system_config, resolve_path
from utils.file_utils import ensure_dir_exists
from utils.logger import get_logger

logger = get_logger(__name__)


def _get_type_config(proposal_type_id: str) -> Dict:
    for t in get_proposal_config()["proposal_types"]:
        if t["id"] == proposal_type_id:
            return t
    raise ValueError(f"Unknown proposal type '{proposal_type_id}'")


def _replace_in_paragraph(paragraph, token: str, value: str) -> None:
    if token not in paragraph.text:
        return
    full_text = paragraph.text.replace(token, value)
    for run in paragraph.runs:
        run.text = ""
    if paragraph.runs:
        paragraph.runs[0].text = full_text
    else:
        paragraph.add_run(full_text)


def _replace_placeholders(document: Document, replacements: Dict[str, str]) -> None:
    for token, value in replacements.items():
        for paragraph in document.paragraphs:
            _replace_in_paragraph(paragraph, token, value)
        for table in document.tables:
            for row in table.rows:
                for cell in row.cells:
                    for paragraph in cell.paragraphs:
                        _replace_in_paragraph(paragraph, token, value)


def export_docx(proposal_type_id: str, field_values: Dict[str, str], reference_numbers: Dict[str, str]) -> str:
    """
    Fills the proposal_type_id's template with field_values (post-edit,
    if the user changed anything in the Proposal Preview step) and
    saves the result to generated_proposals/docx/.

    Returns:
        Absolute path to the saved .docx file.
    """
    type_cfg = _get_type_config(proposal_type_id)
    field_mapping = get_field_mapping()["proposal_types"][proposal_type_id]["fields"]

    template_path = resolve_path(type_cfg["template_path"])
    if not os.path.exists(template_path):
        raise FileNotFoundError(
            f"Template not found for '{proposal_type_id}' at '{template_path}'. "
            f"Run scripts/generate_templates.py to (re)generate it."
        )

    document = Document(template_path)

    replacements = {}
    for field in field_mapping:
        token = field["placeholder_token"]
        value = field_values.get(field["field_name"], "NA")
        replacements[token] = value
    replacements["{{generation_date}}"] = datetime.now().strftime("%d-%b-%Y")

    _replace_placeholders(document, replacements)

    ref_hint = (
        reference_numbers.get("po_number") or reference_numbers.get("tender_number")
        or reference_numbers.get("spr_number") or reference_numbers.get("registration_number") or "proposal"
    )
    safe_ref = "".join(c for c in str(ref_hint) if c.isalnum() or c in ("-", "_")) or "proposal"
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    filename = f"{proposal_type_id}_{safe_ref}_{timestamp}.docx"

    output_dir = resolve_path(get_system_config()["proposal_output"]["docx_output_dir"])
    ensure_dir_exists(output_dir)
    output_path = os.path.join(output_dir, filename)

    document.save(output_path)
    logger.info("DOCX proposal exported: '%s'", output_path)
    return output_path
