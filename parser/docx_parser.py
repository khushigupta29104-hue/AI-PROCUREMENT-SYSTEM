"""
docx_parser.py
===============
Extracts text from Word (.docx) files using python-docx: paragraphs
AND table cell text (procurement docs frequently store key data in
tables).
"""

from typing import Dict

import docx

from parser.base_parser import BaseParser
from utils.logger import get_logger
from utils.text_utils import clean_extracted_text

logger = get_logger(__name__)


class DOCXParser(BaseParser):
    def extract_text(self, file_path: str) -> str:
        try:
            document = docx.Document(file_path)
        except Exception as exc:  # noqa: BLE001
            logger.error("Failed to open DOCX '%s': %s", file_path, exc)
            raise ValueError(f"Could not open DOCX file: {exc}") from exc

        text_parts = []
        for para in document.paragraphs:
            if para.text.strip():
                text_parts.append(para.text.strip())

        for table in document.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                if any(cells):
                    text_parts.append(" | ".join(cells))

        raw_text = "\n".join(text_parts)
        cleaned = clean_extracted_text(raw_text)
        logger.info("Extracted %d characters from DOCX '%s' (%d paragraphs, %d tables)",
                    len(cleaned), file_path, len(document.paragraphs), len(document.tables))
        return cleaned

    def extract_metadata(self, file_path: str) -> Dict:
        try:
            document = docx.Document(file_path)
            core_props = document.core_properties
            return {
                "paragraph_count": len(document.paragraphs),
                "table_count": len(document.tables),
                "title": core_props.title or "",
                "author": core_props.author or "",
                "file_type": "docx",
            }
        except Exception as exc:  # noqa: BLE001
            logger.warning("Could not extract metadata from DOCX '%s': %s", file_path, exc)
            return {"paragraph_count": 0, "table_count": 0, "title": "", "author": "", "file_type": "docx"}
